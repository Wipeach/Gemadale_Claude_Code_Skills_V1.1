#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
增强版：向 PPTX 添加表格（带大量 debug、单位兼容与自动 clamp）
如果你在 pipeline 中修改参数却无效，先看 pipeline 日志里本脚本输出的 [DEBUG] 信息，确认 pipeline 真正调用了哪个脚本与参数。
"""
from typing import Any, Dict, List, Optional
import os
from datetime import datetime
import inspect

# 尝试导入新的 run 接口；若没有，再尝试老接口
add_table_run = None
add_table_func = None
extract_run = None
extract_func = None

try:
    from utils.pptx_gen_add_table_to_slide import run as add_table_run
    print("[DEBUG] Imported add_table.run (new API).")
except Exception:
    try:
        from utils.pptx_gen_add_table_to_slide import add_table_to_slide as add_table_func
        print("[DEBUG] Imported add_table_to_slide (old API).")
    except Exception:
        add_table_run = None
        add_table_func = None
        print("[DEBUG] No add_table module found at import time.")

try:
    from utils.data_processor_extract_table_data import run as extract_run
    print("[DEBUG] Imported extract.run.")
except Exception:
    try:
        from utils.data_processor_extract_table_data import create_table_data_for_presentation as extract_func
        print("[DEBUG] Imported extract.create_table_data_for_presentation.")
    except Exception:
        extract_run = None
        extract_func = None
        print("[DEBUG] No extract module found at import time.")


def _normalize_extracted_table_data(extracted: Any) -> Optional[List[List[Any]]]:
    if extracted is None:
        return None
    if isinstance(extracted, list):
        return extracted
    if isinstance(extracted, dict):
        for key in ("table_data", "table", "result", "data", "payload", "rows", "details"):
            v = extracted.get(key)
            if isinstance(v, list):
                return v
        for v in extracted.values():
            if isinstance(v, list):
                return v
    return None


EMU_PER_INCH = 914400


def inches_to_emu(val_in_inches: float) -> int:
    return int(round(val_in_inches * EMU_PER_INCH))


def run(
    project_name: str,
    pptx_file_path: Optional[str] = None,
    slide_number: int = 1,
    left_position: float = 8.0,
    top_position: float = 2.0,
    table_width: float = 5.5,
    table_height: float = 3.0,
    timestamp: Optional[str] = None,
    table_data_override: Optional[List[List[Any]]] = None,
) -> Dict[str, Any]:
    """
    Extract table data and add it to a PPTX file.

    left_position, top_position, table_width, table_height 的单位默认为英寸（inch）。
    本脚本会在调用 add_table 时同时把英寸与 EMU 两种表示都放入 task 里以提高兼容性。
    """
    print(f"[DEBUG] run() called with project_name={project_name}, slide={slide_number}, left={left_position}, top={top_position}, width={table_width}, height={table_height}, timestamp={timestamp}")
    if not timestamp:
        timestamp = datetime.now().strftime("%Y%m%d")

    base_dir = os.path.join("resources", "working_data", f"{project_name}_{timestamp}", "processed_data")
    os.makedirs(base_dir, exist_ok=True)  # ensure folder exists

    pptx_file = pptx_file_path or os.path.join(base_dir, f"{project_name}_gemdale_housing_project_template.pptx")
    print(f"[DEBUG] Using pptx_file path: {pptx_file}")

    if not os.path.exists(pptx_file):
        msg = f"Input file not found: {pptx_file}"
        print(f"[ERROR] {msg}")
        return {"success": False, "error": msg, "pptx_file_path": pptx_file}

    # 读取 slide 宽度并检查 left_position 是否越界（自动 clamping）
    try:
        from pptx import Presentation as PptxPresentation
        prs_check = PptxPresentation(pptx_file)
        slide_width_in = prs_check.slide_width / EMU_PER_INCH  # EMU -> 英寸
        print(f"[DEBUG] PPTX slide width (inches): {slide_width_in:.3f}\"")
        # 为安全保留 margin 英寸
        right_margin = 0.2
        max_left = slide_width_in - table_width - right_margin
        if max_left < 0:
            max_left = 0.0
        if left_position > max_left:
            print(f"[WARN] left_position {left_position}\" > max allowed {max_left}\"; clamping to {max_left}\"")
            left_position = max_left
        print(f"[DEBUG] Effective left_position (inches) after clamp: {left_position}\"")
    except Exception as e:
        print(f"[INFO] 无法读取 PPTX 以自动调整 left_position: {e}")

    # Get table_data
    if table_data_override is not None:
        table_data = table_data_override
        print("[DEBUG] Using table_data_override.")
    else:
        if extract_run:
            try:
                extracted = extract_run(project_name=project_name, timestamp=timestamp)
            except TypeError:
                try:
                    extracted = extract_run(project_name=project_name)
                except Exception as e:
                    return {"success": False, "error": f"call extract.run failed: {e}", "pptx_file_path": pptx_file}
        elif extract_func:
            try:
                extracted = extract_func(project_name=project_name, timestamp=timestamp)
            except TypeError:
                try:
                    extracted = extract_func(project_name=project_name)
                except Exception as e:
                    return {"success": False, "error": f"call extract_func failed: {e}", "pptx_file_path": pptx_file}
        else:
            return {"success": False, "error": "no extract module available", "pptx_file_path": pptx_file}

        table_data = _normalize_extracted_table_data(extracted)
        if table_data is None:
            return {
                "success": False,
                "error": f"could not parse table data from extractor output: {type(extracted)}",
                "extract_output_preview": (extracted if isinstance(extracted, (str, dict, list)) else str(type(extracted))),
                "pptx_file_path": pptx_file,
            }

    if not isinstance(table_data, list) or len(table_data) == 0 or not isinstance(table_data[0], list):
        return {"success": False, "error": "table_data must be a non-empty list of lists", "pptx_file_path": pptx_file}

    # 构造 task，同时包含英寸与 EMU 表示（提高兼容性）
    task = {
        "slide_number": slide_number,
        # 英寸表示（保留原字段名以兼容旧实现）
        "left_position": left_position,
        "top_position": top_position,
        "table_width": table_width,
        "table_height": table_height,
        # 备用/冗余字段（明确单位）
        "left_in": left_position,
        "top_in": top_position,
        "width_in": table_width,
        "height_in": table_height,
        # EMU（pptx 内部单位）
        "left_emu": inches_to_emu(left_position),
        "top_emu": inches_to_emu(top_position),
        "width_emu": inches_to_emu(table_width),
        "height_emu": inches_to_emu(table_height),
        # 源数据与样式
        "table_data": table_data,
        "style": None,
    }

    print("[DEBUG] Constructed table task (summary):")
    print(f"        slide: {task['slide_number']}, left_in={task['left_in']}\", left_emu={task['left_emu']}, width_in={task['width_in']}\", width_emu={task['width_emu']}")
    print(f"        rows={len(table_data)}, cols={(len(table_data[0]) if len(table_data)>0 else 0)}")

    # 调用新的 run-based add_table 接口（如果存在）
    if add_table_run:
        try:
            print("[DEBUG] Calling add_table.run with table_tasks (includes both inch & EMU values).")
            add_res = add_table_run(project_name=project_name, pptx_file_path=pptx_file, table_tasks=[task], timestamp=timestamp)
            final_pptx = pptx_file
            if isinstance(add_res, dict):
                for candidate_key in ("pptx_output_path", "pptx_file_path", "pptx_output", "output_path"):
                    if candidate_key in add_res and add_res[candidate_key]:
                        final_pptx = add_res[candidate_key]
                        break
                success = bool(add_res.get("success", True))
                print(f"[DEBUG] add_table.run returned dict with success={success}.")
                return {"success": success, "pptx_file_path": final_pptx, "add_table_result": add_res}
            elif isinstance(add_res, bool):
                print(f"[DEBUG] add_table.run returned bool: {add_res}")
                return {"success": add_res, "pptx_file_path": final_pptx}
            else:
                print("[DEBUG] add_table.run returned unknown type; assuming success.")
                return {"success": True, "pptx_file_path": final_pptx, "add_table_raw_result": add_res}
        except Exception as e:
            print(f"[ERROR] calling add_table.run failed: {e}")
            return {"success": False, "error": f"calling add_table.run failed: {e}", "pptx_file_path": pptx_file}

    # 否则尝试旧接口 add_table_func，并尝试传入不同字段以兼容各种实现
    elif add_table_func:
        try:
            sig = None
            try:
                sig = inspect.signature(add_table_func)
                param_names = set(sig.parameters.keys())
            except Exception:
                param_names = set()

            # 基本 kwargs（旧实现最可能接受的）
            kwargs = {
                "pptx_file_path": pptx_file,
                "slide_number": slide_number,
                "left_position": left_position,
                "top_position": top_position,
                "table_data": table_data,
                "table_width": table_width,
                "table_height": table_height,
                "style": None,
            }

            # 如果老函数接受 EMU 字段名字，则加上（提高兼容性）
            if 'left_emu' in param_names or 'left' in param_names:
                kwargs['left_emu'] = inches_to_emu(left_position)
            if 'top_emu' in param_names or 'top' in param_names:
                kwargs['top_emu'] = inches_to_emu(top_position)
            if 'width_emu' in param_names or 'width' in param_names:
                kwargs['width_emu'] = inches_to_emu(table_width)
            if 'height_emu' in param_names:
                kwargs['height_emu'] = inches_to_emu(table_height)

            print(f"[DEBUG] Calling add_table_func with kwargs keys: {list(kwargs.keys())}")
            ok = add_table_func(**kwargs)
            print(f"[DEBUG] add_table_func returned: {ok}")
            return {"success": bool(ok), "pptx_file_path": pptx_file}
        except TypeError as e:
            # 如果签名不匹配，退回到最基础调用（尽量兼容）
            try:
                print(f"[WARN] add_table_func signature mismatch ({e}), trying fallback call without EMU kwargs.")
                ok = add_table_func(pptx_file, slide_number, left_position, top_position, table_data, table_width, table_height, None)
                return {"success": bool(ok), "pptx_file_path": pptx_file}
            except Exception as e2:
                print(f"[ERROR] fallback call to add_table_func also failed: {e2}")
                return {"success": False, "error": f"calling add_table_func failed: {e2}", "pptx_file_path": pptx_file}
        except Exception as e:
            print(f"[ERROR] calling add_table_func failed: {e}")
            return {"success": False, "error": f"calling add_table_func failed: {e}", "pptx_file_path": pptx_file}
    else:
        print("[ERROR] No add_table module available to place table.")
        return {"success": False, "error": "no add_table module available", "pptx_file_path": pptx_file}


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Add project data table to PPTX.")
    parser.add_argument("project_name", help="Project name.")
    parser.add_argument("--pptx", dest="pptx_file", help="PPTX path (optional).")
    parser.add_argument("--slide", dest="slide", type=int, default=1)
    parser.add_argument("--left", dest="left", type=float, default=7.0)  # 更安全的默认值
    parser.add_argument("--top", dest="top", type=float, default=2.0)
    parser.add_argument("--width", dest="width", type=float, default=5.5)
    parser.add_argument("--height", dest="height", type=float, default=3.0)
    parser.add_argument("--timestamp", dest="timestamp", type=str, default=None, help="Optional YYYYMMDD timestamp (default today)")
    args = parser.parse_args()

    out = run(
        project_name=args.project_name,
        pptx_file_path=args.pptx_file,
        slide_number=args.slide,
        left_position=args.left,
        top_position=args.top,
        table_width=args.width,
        table_height=args.height,
        timestamp=args.timestamp
    )
    print("[DEBUG] Final run() output:", out)
